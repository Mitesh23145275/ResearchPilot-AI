from io import BytesIO
from pathlib import Path
import tempfile
import re

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from docx import Document
from docx.shared import Pt
from docx2pdf import convert

from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

router = APIRouter(prefix="/pdf", tags=["PDF Export"])


class PDFRequest(BaseModel):
    topic: str
    draft: str


def title_case(text):
    small = {"of", "and", "the", "for", "to", "in", "on", "with", "by", "at", "from"}

    words = text.strip().split()
    result = []

    for i, w in enumerate(words):
        if w.upper() == "AI":
            result.append("AI")
        elif i > 0 and w.lower() in small:
            result.append(w.lower())
        else:
            result.append(w.capitalize())

    return " ".join(result)


def parse_sections(markdown):
    sections = {
        "ABSTRACT": "",
        "KEYWORDS": "",
        "INTRODUCTION": "",
        "WEB_RESEARCH": "",
        "LITERATURE": "",
        "TECHNICAL": "",
        "INDUSTRY": "",
        "CONCLUSION": "",
        "REFERENCES": "",
    }

    current = None

    for line in markdown.splitlines():
        text = line.strip()

        if not text:
            if current:
                sections[current] += "\n"
            continue

        if text.startswith("# "):
            continue

        if text.startswith("## "):
            heading = text[3:].strip().upper()

            if heading == "ABSTRACT":
                current = "ABSTRACT"
            elif heading == "KEYWORDS":
                current = "KEYWORDS"
            elif "INTRODUCTION" in heading:
                current = "INTRODUCTION"
            elif "WEB RESEARCH" in heading:
                current = "WEB_RESEARCH"
            elif "LITERATURE" in heading:
                current = "LITERATURE"
            elif "TECHNICAL" in heading:
                current = "TECHNICAL"
            elif "INDUSTRY" in heading:
                current = "INDUSTRY"
            elif "CONCLUSION" in heading:
                current = "CONCLUSION"
            elif "REFERENCE" in heading:
                current = "REFERENCES"

            continue

        if current:
            sections[current] += clean_markdown(text) + "\n"

    return {k: v.strip() for k, v in sections.items()}

def clean_markdown(text: str):
    """
    Convert Markdown into clean IEEE text.
    """
    lines = []

    for line in text.splitlines():

        line = line.strip()

        if not line:
            lines.append("")
            continue

        # #### Heading
        line = re.sub(r"^#{3,4}\s*", "", line)

        # remove **bold**
        line = line.replace("**", "")

        # * Heading -> Heading
        if line.startswith("* "):
            line = line[2:]

        # + item -> • item
        if line.startswith("+ "):
            line = "• " + line[2:]

        # - item -> • item
        if line.startswith("- "):
            line = "• " + line[2:]

        lines.append(line)

    return "\n".join(lines)


def replace_placeholder(doc, placeholder, value):
    value = value or ""

    def insert_after(paragraph, text):
        new_p = OxmlElement("w:p")
        paragraph._p.addnext(new_p)
        new_para = Paragraph(new_p, paragraph._parent)

        new_para.style = paragraph.style
        new_para.add_run(text)

        # IEEE alignment
        if text.startswith("•"):
            new_para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            new_para.paragraph_format.left_indent = Pt(18)
            new_para.paragraph_format.first_line_indent = Pt(-9)

        elif re.match(r"^\d+\s", text):
            # Numbered subheadings like "4 Key Findings"
            new_para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            new_para.runs[0].bold = True

        elif text.startswith("["):
            # References
            new_para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

        else:
            # Normal body paragraphs
            new_para.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY

        return new_para

    # Normal paragraphs
    for paragraph in doc.paragraphs:
        if placeholder in paragraph.text:

            # Single-line placeholders: keep the original paragraph
            if placeholder in {
                "{{TITLE}}",
                "{{AUTHOR}}",
                "{{DEPARTMENT}}",
                "{{COLLEGE}}",
                "{{LOCATION}}",
                "{{EMAIL}}",
            }:
                paragraph.text = paragraph.text.replace(placeholder, value)
                return

            # Multi-line content placeholders
            current = paragraph
            current.text = current.text.replace(placeholder, "")

            for line in value.splitlines():
                line = line.strip()

                if not line:
                    continue

                current = insert_after(current, line)

            return

    # Table cells
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if placeholder in paragraph.text:

                        current = paragraph
                        current.text = current.text.replace(placeholder, "")

                        for line in value.splitlines():
                            line = line.strip()

                            if not line:
                                continue

                            current = insert_after(current, line)

                        return


@router.post("/download")
async def download_pdf(request: PDFRequest):
    template = Path(__file__).parent.parent / "templates" / "IEEE_Template.docx"

    if not template.exists():
        raise HTTPException(status_code=500, detail="IEEE_Template.docx not found.")

    doc = Document(template)

    title = title_case(request.topic)
    sections = parse_sections(request.draft)

    replacements = {
        "{{TITLE}}": title,
        "{{AUTHOR}}": "Mitesh Pawar",
        "{{DEPARTMENT}}": "B.Tech AI & Data Science",
        "{{COLLEGE}}": "Your College Name",
        "{{LOCATION}}": "Mumbai, India",
        "{{EMAIL}}": "student@example.com",
        "{{ABSTRACT}}": sections["ABSTRACT"],
        "{{KEYWORDS}}": sections["KEYWORDS"],
        "{{INTRODUCTION}}": sections["INTRODUCTION"],
        "{{WEB_RESEARCH}}": sections["WEB_RESEARCH"],
        "{{LITERATURE}}": sections["LITERATURE"],
        "{{TECHNICAL}}": sections["TECHNICAL"],
        "{{INDUSTRY}}": sections["INDUSTRY"],
        "{{CONCLUSION}}": sections["CONCLUSION"],
        "{{REFERENCES}}": sections["REFERENCES"],
    }

    for key, value in replacements.items():
        replace_placeholder(doc, key, value)

    with tempfile.TemporaryDirectory() as tmp:
        docx_path = Path(tmp) / "paper.docx"
        pdf_path = Path(tmp) / "paper.pdf"

        doc.save(docx_path)

        convert(str(docx_path), str(pdf_path))

        pdf_bytes = pdf_path.read_bytes()

    filename = f"{title.replace(' ', '_')}_research.pdf"

    return StreamingResponse(
        BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )